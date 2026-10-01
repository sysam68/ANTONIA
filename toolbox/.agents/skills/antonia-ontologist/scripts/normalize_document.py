#!/usr/bin/env python3
"""Normalize supported local documents to UTF-8 text for OntoGPT."""

from __future__ import annotations

import argparse
from pathlib import Path


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as document:
            return "\n\n".join(page.get_text("text") for page in document)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(blocks)
    raise ValueError(f"unsupported document format: {suffix or '(none)'}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if not args.input.is_file():
        raise SystemExit(f"input document not found: {args.input}")
    text = extract(args.input).strip()
    if not text:
        raise SystemExit("document contains no extractable text; OCR is not supported")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(text + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
